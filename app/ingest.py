"""Read PO workbooks without assuming sheet names or column positions.
Extraction rules are v1's (docs/manufaktur-queue-design.md, Data prep)."""
import datetime
import openpyxl

SCAN_ROWS = 15
FIELD_TOKENS = {"size": "size", "berat/pcs": "berat_per_pcs", "berat": "berat_total", "no": "po_no"}


def load(path):
    return openpyxl.load_workbook(path, data_only=True)


def _token(value):
    parts = str(value).split() if value is not None else []
    tok = parts[0].lower() if parts else ""
    return "thickness" if tok.startswith(("tick", "thick")) else tok


def _columns(ws, row):
    cols = {}
    for c in range(1, ws.max_column + 1):
        tok = _token(ws.cell(row=row, column=c).value)
        name = "thickness" if tok == "thickness" else FIELD_TOKENS.get(tok)
        if name and name not in cols:
            cols[name] = c
    return cols


def _machines(ws, row):
    out = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=row, column=c).value
        if isinstance(v, str) and v.strip().upper().startswith("MESIN"):
            out.append((v.strip(), c, c + 1))
    return out


def find_header_row(ws):
    for r in range(1, min(SCAN_ROWS, ws.max_row) + 1):
        cols = _columns(ws, r)
        if {"thickness", "size", "berat_per_pcs", "po_no"} <= cols.keys() and _machines(ws, r + 1):
            return r
    return None


def snapshot_date(ws, header_row):
    for r in range(1, header_row):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, datetime.datetime):
                return v.date().isoformat()
    return ""


def scan_workbook(wb):
    found = []
    for ws in wb.worksheets:
        h = find_header_row(ws)
        if h:
            found.append({"sheet": ws.title, "header_row": h, "snapshot_date": snapshot_date(ws, h),
                          "n_rows": len(extract_sheet(ws, h))})
    return found


def extract_sheet(ws, header_row):
    cols, machines = _columns(ws, header_row), _machines(ws, header_row + 1)
    mold_row, rows, thickness = header_row + 2, [], None
    for r in range(header_row + 3, ws.max_row + 1):
        t = ws.cell(row=r, column=cols["thickness"]).value
        if isinstance(t, str):
            continue                      # TOTAL / GRAND TOTAL
        if t is not None:
            thickness = float(t)
        po_no = ws.cell(row=r, column=cols["po_no"]).value
        size = ws.cell(row=r, column=cols["size"]).value
        if po_no is None or size is None:
            continue                      # empty slot for this size
        mesin = mesin_size = ""
        for name, ca, cb in machines:
            col = next((c for c in (ca, cb) if ws.cell(row=r, column=c).value is not None), None)
            if col:
                mesin, mesin_size = name, str(ws.cell(row=mold_row, column=col).value).replace("\n", " ")
                break
        rows.append({
            "po_no": str(po_no).strip(),
            "thickness": thickness,
            "size": str(size).replace("-", " ").strip(),
            "berat_per_pcs": float(ws.cell(row=r, column=cols["berat_per_pcs"]).value or 0),
            "berat_total": float(ws.cell(row=r, column=cols["berat_total"]).value or 0),
            "mesin": mesin, "mesin_size": mesin_size,
        })
    return rows
