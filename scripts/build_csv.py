"""Rebuild 2025.csv / 2026.csv from the raw TISCO workbook, carrying
berat_per_pcs through (the previous flatten dropped it).

Columns are located by header text (row 3 for the base fields, row 4 for
the machine names), so inserting a column in the sheet doesn't silently
shift what gets read.

Run: python scripts/build_csv.py
"""
import csv
import openpyxl

RAW_XLSX = "UPDATE SISA  PO PIPA TISCO TAHUN 2025.xlsx"
HEADER_ROW = 3          # TICKNES / SIZE / BERAT/PCS / Berat / NO PO
MACHINE_ROW = 4         # MESIN 1 .. MESIN 7, each spanning two columns
MOLD_ROW = 5            # the size/mold label under each machine column
FIRST_DATA_ROW = 6
FIELDS = ["po_no", "thickness", "size", "berat_per_pcs", "berat_total", "mesin", "mesin_size"]


def _first_token(value):
    tokens = str(value).split() if value is not None else []
    return tokens[0].lower() if tokens else ""


def find_columns(ws):
    """Map field names to 1-based column indexes by reading the header rows.
    Raises if any expected header is missing."""
    wanted = {"thickness": "thickness", "size": "size", "berat/pcs": "berat_per_pcs",
              "berat": "berat_total", "no": "po_no"}
    cols = {}
    for c in range(1, ws.max_column + 1):
        token = _first_token(ws.cell(row=HEADER_ROW, column=c).value)
        if token.startswith("tick") or token.startswith("thick"):
            token = "thickness"  # sheet spells it "TICKNES"
        if token in wanted and wanted[token] not in cols:
            cols[wanted[token]] = c
    missing = [name for name in wanted.values() if name not in cols]
    if missing:
        raise ValueError(f"{ws.title}: header row {HEADER_ROW} is missing {missing}")

    machines = []
    for c in range(1, ws.max_column + 1):
        value = ws.cell(row=MACHINE_ROW, column=c).value
        if isinstance(value, str) and value.strip().upper().startswith("MESIN"):
            machines.append((value.strip(), c, c + 1))
    if not machines:
        raise ValueError(f"{ws.title}: no MESIN headers found in row {MACHINE_ROW}")
    return cols, machines


def extract_sheet(ws):
    cols, machines = find_columns(ws)
    rows = []
    current_thickness = None
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        thickness_raw = ws.cell(row=r, column=cols["thickness"]).value
        if isinstance(thickness_raw, str):
            continue  # "TOTAL" / "GRAND TOTAL" subtotal row
        if thickness_raw is not None:
            current_thickness = thickness_raw

        po_no = ws.cell(row=r, column=cols["po_no"]).value
        if po_no is None:
            continue  # no order for this size in this section

        size = ws.cell(row=r, column=cols["size"]).value
        if size is None:
            continue
        size = str(size).replace("-", " ").strip()

        mesin, mesin_size = "", ""
        for name, col_a, col_b in machines:
            for col in (col_a, col_b):
                if ws.cell(row=r, column=col).value is not None:
                    mesin = name
                    mesin_size = str(ws.cell(row=MOLD_ROW, column=col).value).replace("\n", " ")
                    break
            if mesin:
                break

        rows.append({
            "po_no": str(po_no),
            "thickness": current_thickness,
            "size": size,
            "berat_per_pcs": ws.cell(row=r, column=cols["berat_per_pcs"]).value,
            "berat_total": ws.cell(row=r, column=cols["berat_total"]).value or 0,
            "mesin": mesin,
            "mesin_size": mesin_size,
        })
    return rows


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def self_check(rows_2025, rows_2026):
    assert len(rows_2025) > 100, f"2025 extraction looks too small: {len(rows_2025)} rows"
    assert len(rows_2026) > 50, f"2026 extraction looks too small: {len(rows_2026)} rows"

    match = [r for r in rows_2026 if r["po_no"] == "PO 001" and r["size"] == '1/2"']
    assert match, "expected a PO 001 1/2\" row in 2026 data"
    assert abs(match[0]["berat_per_pcs"] - 1.31) < 1e-9, f"unexpected berat_per_pcs: {match[0]}"
    assert match[0]["mesin"] == "MESIN 2", f"unexpected mesin: {match[0]}"

    assert not any(r["po_no"].upper() in ("TOTAL", "GRAND TOTAL") for r in rows_2025), \
        "a TOTAL row leaked into extraction"

    print(f"self-check OK: {len(rows_2025)} rows (2025), {len(rows_2026)} rows (2026)")


def main():
    wb = openpyxl.load_workbook(RAW_XLSX, data_only=True)
    rows_2025 = extract_sheet(wb["2025"])
    rows_2026 = extract_sheet(wb["2026"])
    write_csv(rows_2025, "2025.csv")
    write_csv(rows_2026, "2026.csv")
    self_check(rows_2025, rows_2026)


if __name__ == "__main__":
    main()
