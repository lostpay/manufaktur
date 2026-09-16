"""Run: python tests/test_ingest.py"""
import datetime
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from app.ingest import scan_workbook, extract_sheet


def synthetic_workbook():
    """Two matching sheets (one with the header at row 5 and an extra column
    before SIZE), one decoy sheet, a TOTAL row, a blank-PO row, mixed size spellings."""
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Sisa PO Sept 2026"
    ws["G2"] = datetime.datetime(2026, 9, 1)
    ws.append([]); ws.append([])                                      # G2 put the cursor at row 2; rows 3-4 blank, header at 5
    ws.append([None, "TICKNES", "kode", "SIZE", None, "BERAT/PCS", "Qty", "Berat", "NO PO", "ZG40", None, "ZG30", None])
    ws.append([None, None, None, None, None, None, None, None, None, "MESIN 1", None, "MESIN 2", None])
    ws.append([None, None, None, None, None, None, None, None, None, '3/4"', "10X20", '1/2"\n5/8"', "10X10"])
    ws.append([None, 0.8, "x", '1/2"', None, 1.4, 100, 140, "PO 001", None, None, 140, None])
    ws.append([None, None, "x", '1-1/4"', None, 3.0, 10, 30, "PO 001", None, None, None, None])
    ws.append([None, None, "x", '3/4"', None, None, None, None, None, None, None, None, None])   # empty slot
    ws.append([None, None, "x", "20X20", "HL", 2.9, 0, 0, "PO 002", None, None, None, None])
    ws.append([None, "TOTAL", None, None, None, None, 110, 170, None, None, None, None, None])
    ws.append([None, 1.0, "x", '1/2"', None, 1.7, 50, 85, "PO 003", None, None, 85, None])
    other = wb.create_sheet("2027")
    other.append([None, "TICKNES", "SIZE", None, "BERAT/PCS", "Qty", "Berat", "NO PO", "ZG40"])
    other.append([None, None, None, None, None, None, None, None, "MESIN 1"])
    other.append([None, None, None, None, None, None, None, None, '3/4"'])
    other.append([None, 0.8, '3/4"', None, 2.2, 10, 22, "PO 009", 22])
    wb.create_sheet("Notes").append(["just", "text"])
    return wb


def test_scan_finds_only_matching_sheets():
    found = scan_workbook(synthetic_workbook())
    assert [f["sheet"] for f in found] == ["Sisa PO Sept 2026", "2027"]
    assert found[0]["header_row"] == 5 and found[0]["snapshot_date"] == "2026-09-01"
    assert found[0]["n_rows"] == 4 and found[1]["n_rows"] == 1 and found[1]["snapshot_date"] == ""


def test_extract_applies_v1_rules_at_a_detected_header():
    wb = synthetic_workbook()
    rows = extract_sheet(wb["Sisa PO Sept 2026"], 5)
    assert [r["po_no"] for r in rows] == ["PO 001", "PO 001", "PO 002", "PO 003"]   # empty slot + TOTAL skipped
    assert [r["thickness"] for r in rows] == [0.8, 0.8, 0.8, 1.0]                     # carried forward
    assert rows[1]["size"] == '1 1/4"'                                                # dash normalized
    assert rows[2]["size"] == "20X20" and rows[2]["berat_total"] == 0
    assert rows[0]["mesin"] == "MESIN 2" and rows[0]["mesin_size"] == '1/2" 5/8"'
    assert rows[3]["berat_per_pcs"] == 1.7


if __name__ == "__main__":
    tests = [v for k, v in list(globals().items()) if k.startswith("test_")]
    for t in tests:
        t(); print(f"OK: {t.__name__}")
    print(f"\n{len(tests)} tests passed")
